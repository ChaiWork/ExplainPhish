"""
ExplainPhish SOC Incident Report Exporter (PDF & Word DOCX)
============================================================
Compiles comprehensive, enterprise-ready incident response reports
into downloadable PDF and Microsoft Word (.docx) formats.
"""

from __future__ import annotations

import html
import io
import logging
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfgen import canvas
from reportlab.platypus import (
    HRFlowable,
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


class NumberedCanvas(canvas.Canvas):
    """Adds Page X of Y and SOC confidentiality header/footer to each page."""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count: int):
        self.saveState()
        self.setFont("Helvetica-Bold", 8)
        self.setFillColor(colors.HexColor("#64748B"))

        # Running Header
        self.drawString(54, 755, "EXPLAINPHISH AUTONOMOUS SOC WORKSTATION — THREAT DOSSIER")
        self.drawRightString(612 - 54, 755, "STRICTLY CONFIDENTIAL // TLP:AMBER")

        # Top rule
        self.setStrokeColor(colors.HexColor("#CBD5E1"))
        self.setLineWidth(0.5)
        self.line(54, 748, 612 - 54, 748)

        # Bottom rule
        self.line(54, 45, 612 - 54, 45)

        # Running Footer
        self.setFont("Helvetica", 8)
        self.drawString(54, 32, "ExplainPhish Autonomous Incident Response Pipeline v2.0")
        page_str = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(612 - 54, 32, page_str)

        self.restoreState()


logger = logging.getLogger("ExplainPhish.Exporter")


def _safe_xml(text: Any) -> str:
    """
    Escapes text so that it can be safely rendered inside a ReportLab Paragraph.
    Replaces &, <, > with XML entities so paraparser NEVER encounters unclosed tags.
    """
    if text is None:
        return ""
    return html.escape(str(text), quote=False)


def _safe_markdown_to_paragraph_xml(text: str) -> str:
    """
    Safely converts markdown formatting to ReportLab-compliant XML:
    - Escapes all raw &, <, > first to prevent XML paraparser crashes
    - Re-applies safe <b>, <i>, and <font> tags from markdown tokens
    """
    if not text:
        return ""
    # 1. Escape all XML entities first
    s = html.escape(str(text), quote=False)

    # 2. Convert markdown bold **text** to <b>text</b>
    s = re.sub(r"\*\*([^\*\n]+?)\*\*", r"<b>\1</b>", s)

    # 3. Convert markdown italic *text* to <i>text</i>
    s = re.sub(r"(?<!\*)\*([^\*\n]+?)\*(?!\*)", r"<i>\1</i>", s)

    # 4. Convert markdown inline code `text` to Courier font
    s = re.sub(r"`([^`\n]+?)`", r"<font name='Courier' size='7.5'>\1</font>", s)

    return s.strip()


def _make_safe_paragraph(text: Any, style: ParagraphStyle, is_markdown: bool = False) -> Paragraph:
    """
    Constructs a ReportLab Paragraph guaranteed not to crash with XML paraparser errors.
    If any XML paraparser error still occurs, falls back to raw plain-text escaping.
    """
    try:
        content = _safe_markdown_to_paragraph_xml(text) if is_markdown else _safe_xml(text)
        return Paragraph(content, style)
    except Exception:
        raw_escaped = html.escape(str(text or ""), quote=False)
        return Paragraph(raw_escaped, style)


def _clean_markdown_text(text: str) -> str:
    """Strips raw markdown formatting tokens for plain text insertion in Word."""
    if not text:
        return ""
    t = re.sub(r"^###+\s*", "", text)
    t = re.sub(r"\*\*([^*]+)\*\*", r"\1", t)
    t = re.sub(r"\*([^*]+)\*", r"\1", t)
    t = re.sub(r"`([^`]+)`", r"\1", t)
    return t.strip()


# ─────────────────────────────────────────────────────────────────────────────
# 1. PDF Report Generation (ReportLab)
# ─────────────────────────────────────────────────────────────────────────────


def generate_pdf_report(state: Dict[str, Any]) -> bytes:
    """Generates an executive-ready PDF Incident Report as raw bytes."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54,
    )

    styles = getSampleStyleSheet()

    # Custom Palette & Styles
    c_navy = colors.HexColor("#0F172A")
    c_slate = colors.HexColor("#334155")
    c_blue = colors.HexColor("#0284C7")
    c_gray_bg = colors.HexColor("#F8FAFC")
    c_border = colors.HexColor("#E2E8F0")

    verdict = state.get("final_verdict") or state.get("ensemble_verdict", "UNKNOWN")
    threat = state.get("threat_level", "UNKNOWN")
    is_malicious = "MALICIOUS" in verdict
    is_suspicious = "SUSPICIOUS" in verdict

    badge_bg = colors.HexColor("#DC2626") if is_malicious else (colors.HexColor("#EA580C") if is_suspicious else colors.HexColor("#059669"))

    style_title = ParagraphStyle(
        "DocTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        textColor=c_navy,
        spaceAfter=4,
    )

    style_subtitle = ParagraphStyle(
        "DocSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=12,
        textColor=c_slate,
        spaceAfter=12,
    )

    style_h2 = ParagraphStyle(
        "Heading2",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=16,
        textColor=c_navy,
        spaceBefore=14,
        spaceAfter=6,
        keepWithNext=True,
    )

    style_h3 = ParagraphStyle(
        "Heading3",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=10,
        leading=14,
        textColor=c_navy,
        spaceBefore=8,
        spaceAfter=3,
        keepWithNext=True,
    )

    style_body = ParagraphStyle(
        "Body",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#1E293B"),
        spaceAfter=5,
    )

    style_bullet = ParagraphStyle(
        "BulletItem",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#1E293B"),
        leftIndent=14,
        spaceAfter=3,
    )

    style_cell = ParagraphStyle(
        "Cell",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8,
        leading=11,
        textColor=colors.HexColor("#1E293B"),
    )

    style_cell_bold = ParagraphStyle(
        "CellBold",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8,
        leading=11,
        textColor=colors.HexColor("#0F172A"),
    )

    style_cell_header = ParagraphStyle(
        "CellHeader",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8,
        leading=11,
        textColor=colors.white,
    )

    story = []
    timestamp = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
    filename = state.get("file_name", "artifact")
    sha256 = state.get("file_hash_sha256", "N/A")
    md5 = state.get("file_hash_md5", "N/A")
    conf = state.get("confidence_score", 0.0)

    # Document Title & Classification Header
    story.append(Paragraph(f"SOC Incident Response Report — {_safe_xml(filename)}", style_title))
    story.append(Paragraph(f"Generated: {timestamp} | Classification: <b>{_safe_xml(threat)}</b> | Pipeline: ExplainPhish v2.0", style_subtitle))

    # Executive Verdict Banner Table
    banner_data = [
        [
            Paragraph(f"<font color='white'><b>VERDICT: {_safe_xml(verdict.upper())}</b></font>", style_cell_header),
            Paragraph(f"<font color='white'><b>THREAT LEVEL: {_safe_xml(threat)}</b></font>", style_cell_header),
            Paragraph(f"<font color='white'><b>CONFIDENCE: {conf:.1%}</b></font>", style_cell_header),
        ]
    ]
    t_banner = Table(banner_data, colWidths=[200, 150, 154])
    t_banner.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), badge_bg),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(t_banner)
    story.append(Spacer(1, 10))

    # 1. Executive Summary Table
    story.append(Paragraph("1. Executive Summary", style_h2))
    size_str = f"{state.get('file_size_bytes', 0):,} bytes"
    badge_hex = "#DC2626" if is_malicious else ("#EA580C" if is_suspicious else "#059669")
    summary_data = [
        [Paragraph("File Name", style_cell_bold), _make_safe_paragraph(str(filename), style_cell)],
        [Paragraph("Verdict", style_cell_bold), Paragraph(f"<b><font color='{badge_hex}'>{_safe_xml(verdict.upper())}</font></b>", style_cell)],
        [Paragraph("Threat Level", style_cell_bold), _make_safe_paragraph(str(threat), style_cell_bold)],
        [Paragraph("Ensemble Confidence", style_cell_bold), _make_safe_paragraph(f"{conf:.1%} ({state.get('confidence_band', 'UNKNOWN')} confidence band)", style_cell)],
        [Paragraph("Format & Size", style_cell_bold), _make_safe_paragraph(f"{state.get('format_display', state.get('file_format', 'Unknown'))} | {size_str}", style_cell)],
        [Paragraph("SHA-256", style_cell_bold), Paragraph(f"<font name='Courier' size=6.5>{_safe_xml(sha256)}</font>", style_cell)],
        [Paragraph("MD5", style_cell_bold), Paragraph(f"<font name='Courier' size=6.5>{_safe_xml(md5)}</font>", style_cell)],
    ]
    if state.get("source_url"):
        summary_data.append([Paragraph("Source Target URL", style_cell_bold), _make_safe_paragraph(str(state["source_url"]), style_cell)])

    t_summary = Table(summary_data, colWidths=[140, 364])
    t_summary.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), c_gray_bg),
        ("GRID", (0, 0), (-1, -1), 0.5, c_border),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(t_summary)
    story.append(Spacer(1, 6))

    # Investigation Pipeline Rules Applied
    rules = state.get("applied_rules") or {}
    if rules:
        story.append(Paragraph("Investigation Pipeline Rules Applied", style_h3))
        rules_data = [[
            Paragraph("Stage", style_cell_header),
            Paragraph("Trigger & Operational Rule Evaluation", style_cell_header),
        ]]
        stage_order = [
            "intake_safety", "feature_extraction", "ml_ensemble",
            "explainability", "deep_threat_analysis", "mitre_mapping",
            "analyst_review", "cognitive_triage", "soc_report"
        ]
        for sk in stage_order:
            if sk in rules:
                stitle = sk.replace("_", " ").title()
                rules_data.append([
                    _make_safe_paragraph(stitle, style_cell_bold),
                    _make_safe_paragraph(str(rules[sk]), style_cell),
                ])
        t_rules = Table(rules_data, colWidths=[140, 364])
        t_rules.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), c_navy),
            ("GRID", (0, 0), (-1, -1), 0.5, c_border),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, c_gray_bg]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        story.append(t_rules)
        story.append(Spacer(1, 6))

    # Autonomous Decision Rationale
    deep = state.get("deep_analysis") or {}
    if deep.get("rationale"):
        story.append(Paragraph(f"<b>Autonomous Forensics Decision:</b> {_safe_xml(deep['rationale'])}", style_body))
        story.append(Spacer(1, 6))

    # 2. AI Threat Intelligence Narrative
    llm = state.get("llm_triage") or {}
    raw_narrative = llm.get("executive_threat_narrative") or state.get("llm_narrative", "")
    if raw_narrative:
        story.append(Paragraph("2. AI Threat Intelligence & Cognitive Forensics", style_h2))
        provider_name = _safe_xml(llm.get("provider", "ExplainPhish AI"))
        model_name = _safe_xml(llm.get("model", "Cognitive Reasoner"))
        status_badge = _safe_xml(llm.get("status_badge", ""))
        if status_badge:
            sub_desc = f"{status_badge} (Engine: <b>{provider_name}</b> | Model: {model_name})"
        else:
            sub_desc = f"Analysis synthesized by <b>{provider_name}</b> ({model_name})"
        story.append(Paragraph(f"<i>{sub_desc}</i>", style_subtitle))

        lines = raw_narrative.split("\n")
        for raw_line in lines:
            line = raw_line.strip()
            if not line:
                continue
            if line.startswith("#"):
                clean_heading = re.sub(r"^#+\s*", "", line)
                story.append(_make_safe_paragraph(f"<b>{_safe_xml(clean_heading)}</b>", style_h3, is_markdown=False))
                continue
            if re.match(r"^(\*|-|\d+\.)\s+", line):
                m_num = re.match(r"^(\d+\.)\s*", line)
                cleaned_item = re.sub(r"^(\*|-|\d+\.)\s*", "", line)
                prefix = f"{m_num.group(1)} " if m_num else "• "
                story.append(_make_safe_paragraph(f"{prefix}{cleaned_item}", style_bullet, is_markdown=True))
                continue
            story.append(_make_safe_paragraph(line, style_body, is_markdown=True))

        story.append(Spacer(1, 6))

    # 3. Model Consensus Table
    preds = state.get("model_predictions", [])
    if preds:
        story.append(Paragraph("3. Multi-Model ML Ensemble Consensus", style_h2))
        story.append(Paragraph(f"Consensus voting evaluated across {len(preds)} standardized architectures with 100% feature parity.", style_body))
        pred_table_data = [[
            Paragraph("Architecture", style_cell_header),
            Paragraph("Prediction", style_cell_header),
            Paragraph("Phishing Probability", style_cell_header),
            Paragraph("Status", style_cell_header),
        ]]
        for p in preds:
            pl = "MALICIOUS" if p.get("prediction") == 1 else "BENIGN"
            prob_val = p.get("probability", 0.0)
            icon = "⚠️" if p.get("prediction") == 1 else "✅"
            pred_table_data.append([
                _make_safe_paragraph(str(p.get("model", "-")), style_cell_bold),
                _make_safe_paragraph(pl, style_cell),
                _make_safe_paragraph(f"{prob_val:.1%}", style_cell),
                _make_safe_paragraph(icon, style_cell),
            ])
        t_preds = Table(pred_table_data, colWidths=[180, 120, 134, 70])
        t_preds.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), c_navy),
            ("GRID", (0, 0), (-1, -1), 0.5, c_border),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, c_gray_bg]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(t_preds)
        consensus_text = "Unanimous agreement across all models." if state.get("unanimous") else "Split consensus vote resolved through LangGraph deep analysis."
        story.append(Spacer(1, 4))
        story.append(Paragraph(f"<b>Consensus Status:</b> {consensus_text}", style_body))
        story.append(Spacer(1, 6))

    # 4. Explainable AI Decision Drivers
    drivers = state.get("top_risk_drivers", [])
    if drivers:
        story.append(Paragraph("4. Explainable AI (XAI) — Top Decision Drivers", style_h2))
        story.append(Paragraph("Feature impact on model classification derived from normalized feature importance and standardized training deviations (<i>z = (x - &mu;) / &sigma;</i>):", style_body))
        xai_data = [[
            Paragraph("Feature", style_cell_header),
            Paragraph("Operational Description", style_cell_header),
            Paragraph("Raw Value", style_cell_header),
            Paragraph("Z-Score", style_cell_header),
            Paragraph("Direction", style_cell_header),
            Paragraph("Risk Impact", style_cell_header),
        ]]
        for d in drivers[:6]:
            is_risk = d.get("direction") == "↑"
            dir_text = "🔺 Increases Risk" if is_risk else "🔻 Reduces Risk"
            feat_name = _safe_xml(str(d.get("feature", "-")))
            z_score_str = f"{d.get('std_value', 0.0):+.2f}"
            impact_str = f"({d.get('impact', 0.0):+.4f})"
            xai_data.append([
                Paragraph(f"<font name='Courier'><b>{feat_name}</b></font>", style_cell_bold),
                _make_safe_paragraph(str(d.get("description", "-")), style_cell),
                _make_safe_paragraph(str(d.get("raw_value", "-")), style_cell),
                _make_safe_paragraph(z_score_str, style_cell),
                _make_safe_paragraph(str(d.get("direction", "")), style_cell),
                _make_safe_paragraph(f"{dir_text} {impact_str}", style_cell),
            ])
        t_xai = Table(xai_data, colWidths=[90, 174, 60, 50, 45, 85])
        t_xai.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), c_navy),
            ("GRID", (0, 0), (-1, -1), 0.5, c_border),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, c_gray_bg]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(t_xai)
        story.append(Spacer(1, 6))

    # 5. Deep Forensic Static Telemetry & IOCs
    indicators = deep.get("indicators", [])
    iocs = state.get("indicators_of_compromise", [])
    if indicators or iocs:
        story.append(Paragraph("5. Deep Forensic Static Telemetry", style_h2))
        if indicators:
            story.append(Paragraph("Autonomous static inspection findings:", style_body))
            for ind in indicators:
                story.append(Paragraph(f"• {_safe_xml(ind)}", style_bullet))
            story.append(Spacer(1, 4))

        if iocs:
            story.append(Paragraph("Extracted Indicators of Compromise (IOCs)", style_h3))
            ioc_data = [[
                Paragraph("Indicator Category", style_cell_header),
                Paragraph("Extracted Value / Target URI", style_cell_header),
            ]]
            for ioc in iocs[:8]:
                ioc_val = _safe_xml(str(ioc.get("value", "-")))
                ioc_data.append([
                    _make_safe_paragraph(str(ioc.get("type", "-")), style_cell_bold),
                    Paragraph(f"<font size=7 name='Courier'>{ioc_val}</font>", style_cell),
                ])
            t_ioc = Table(ioc_data, colWidths=[150, 354])
            t_ioc.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), c_navy),
                ("GRID", (0, 0), (-1, -1), 0.5, c_border),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, c_gray_bg]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.append(t_ioc)
            story.append(Spacer(1, 6))

    # 6. MITRE ATT&CK Mapping
    mitre = state.get("mitre_tactics", [])
    if mitre:
        story.append(Paragraph("6. MITRE ATT&CK Mapping", style_h2))
        mitre_data = [[
            Paragraph("Technique ID", style_cell_header),
            Paragraph("Tactic", style_cell_header),
            Paragraph("Technique Name", style_cell_header),
            Paragraph("Operational Description", style_cell_header),
        ]]
        for m in mitre:
            mitre_data.append([
                Paragraph(f"<font color='#0284C7' name='Courier'><b>{_safe_xml(m.get('id', '-'))}</b></font>", style_cell),
                _make_safe_paragraph(str(m.get("tactic", "-")), style_cell),
                _make_safe_paragraph(str(m.get("technique", "-")), style_cell_bold),
                _make_safe_paragraph(str(m.get("description", "-")), style_cell),
            ])
        t_mitre = Table(mitre_data, colWidths=[80, 100, 140, 184])
        t_mitre.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), c_navy),
            ("GRID", (0, 0), (-1, -1), 0.5, c_border),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, c_gray_bg]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(t_mitre)
        story.append(Spacer(1, 6))

    # 7. Prescriptive SOAR Incident Response Playbook
    playbook = state.get("soc_playbook_actions", [])
    if playbook:
        story.append(Paragraph("7. Prescriptive SOAR Incident Response Playbook", style_h2))
        soar_data = [[
            Paragraph("Stage", style_cell_header),
            Paragraph("Response Tier", style_cell_header),
            Paragraph("Prescribed Operational Mitigation", style_cell_header),
        ]]
        for a in playbook:
            tier_val = _safe_xml(str(a.get("tier", "-")))
            soar_data.append([
                _make_safe_paragraph(str(a.get("stage", "-")), style_cell_bold),
                Paragraph(f"<font color='#0284C7'><b>{tier_val}</b></font>", style_cell),
                _make_safe_paragraph(str(a.get("action", "-")), style_cell),
            ])
        t_soar = Table(soar_data, colWidths=[100, 120, 284])
        t_soar.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), c_navy),
            ("GRID", (0, 0), (-1, -1), 0.5, c_border),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, c_gray_bg]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(t_soar)
        story.append(Spacer(1, 6))

    # 8. Human-in-the-Loop (HITL) Analyst Verification
    analyst_review = state.get("analyst_review") or {}
    if analyst_review:
        story.append(Paragraph("8. Human-in-the-Loop (HITL) Analyst Verification", style_h2))
        hitl_data = [
            [Paragraph("Review Status", style_cell_bold), _make_safe_paragraph(str(analyst_review.get("status", "N/A")), style_cell)],
            [Paragraph("Analyst Decision", style_cell_bold), _make_safe_paragraph(str(analyst_review.get("decision", "APPROVE")), style_cell_bold)],
            [Paragraph("Signed-off By", style_cell_bold), _make_safe_paragraph(str(analyst_review.get("analyst_id", "SOC-Analyst")), style_cell)],
            [Paragraph("Review Timestamp", style_cell_bold), _make_safe_paragraph(str(analyst_review.get("reviewed_at", timestamp)), style_cell)],
            [Paragraph("Analyst Justification", style_cell_bold), _make_safe_paragraph(str(analyst_review.get("notes", "Standard verification sign-off.")), style_cell)],
        ]
        t_hitl = Table(hitl_data, colWidths=[140, 364])
        t_hitl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), c_gray_bg),
            ("GRID", (0, 0), (-1, -1), 0.5, c_border),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(t_hitl)
        story.append(Spacer(1, 6))

    story.append(Spacer(1, 8))
    story.append(HRFlowable(width="100%", thickness=0.5, color=c_border, spaceBefore=4, spaceAfter=8))
    story.append(Paragraph("<i>Report automatically generated by ExplainPhish Autonomous Security Agent. For SOC escalations, submit hash to SIEM.</i>", style_subtitle))

    try:
        doc.build(story, canvasmaker=NumberedCanvas)
    except Exception as exc:
        logger.warning(f"ReportLab primary build encountered error: {exc}. Retrying with strict plain text fallback.")
        fallback_buffer = io.BytesIO()
        fallback_doc = SimpleDocTemplate(fallback_buffer, pagesize=letter, leftMargin=54, rightMargin=54, topMargin=54, bottomMargin=54)
        safe_story = []
        for flowable in story:
            if isinstance(flowable, Paragraph):
                plain_txt = re.sub(r"<[^>]+>", "", flowable.text)
                safe_story.append(Paragraph(_safe_xml(plain_txt), flowable.style))
            else:
                safe_story.append(flowable)
        fallback_doc.build(safe_story, canvasmaker=NumberedCanvas)
        return fallback_buffer.getvalue()

    return buffer.getvalue()


# ─────────────────────────────────────────────────────────────────────────────
# 2. Microsoft Word (.docx) Report Generation (python-docx)
# ─────────────────────────────────────────────────────────────────────────────


def generate_docx_report(state: Dict[str, Any]) -> bytes:
    """Generates an executive-ready Microsoft Word (.docx) Incident Report as raw bytes."""
    doc = docx.Document()

    # Set 0.75 inch margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    verdict = state.get("final_verdict") or state.get("ensemble_verdict", "UNKNOWN")
    threat = state.get("threat_level", "UNKNOWN")
    conf = state.get("confidence_score", 0.0)
    band = state.get("confidence_band", "UNKNOWN")
    filename = state.get("file_name", "artifact")
    sha256 = state.get("file_hash_sha256", "N/A")
    md5 = state.get("file_hash_md5", "N/A")
    timestamp = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")

    # Document Header Title
    title = doc.add_paragraph()
    r_title = title.add_run(f"SOC Incident Response Report — {filename}")
    r_title.bold = True
    r_title.font.size = Pt(18)
    r_title.font.color.rgb = RGBColor(15, 23, 42)  # Navy
    title.paragraph_format.space_after = Pt(2)

    sub = doc.add_paragraph()
    r_sub = sub.add_run(f"Generated: {timestamp} | Classification: {threat} | Pipeline: ExplainPhish Autonomous SOC Agent v2.0")
    r_sub.font.size = Pt(9)
    r_sub.font.color.rgb = RGBColor(100, 116, 139)
    sub.paragraph_format.space_after = Pt(12)

    # 1. Executive Summary Table
    h1 = doc.add_heading("1. Executive Summary", level=2)
    h1.paragraph_format.space_before = Pt(8)
    h1.paragraph_format.space_after = Pt(4)

    summary_rows = [
        ("File Name", str(filename)),
        ("Verdict", str(verdict)),
        ("Threat Level", str(threat)),
        ("Ensemble Confidence", f"{conf:.1%} ({band} confidence band)"),
        ("Format & Size", f"{state.get('format_display', state.get('file_format', 'Unknown'))} | {state.get('file_size_bytes', 0):,} bytes"),
        ("SHA-256", sha256),
        ("MD5", md5),
    ]
    if state.get("source_url"):
        summary_rows.append(("Source Target URL", str(state["source_url"])))

    t_sum = doc.add_table(rows=len(summary_rows), cols=2)
    t_sum.style = "Table Grid"
    t_sum.alignment = WD_TABLE_ALIGNMENT.CENTER

    for idx, (lbl, val) in enumerate(summary_rows):
        r = t_sum.rows[idx]
        _set_cell_text(r.cells[0], lbl, bold=True, bg_color="F1F5F9")
        _set_cell_text(r.cells[1], val, bold=(idx == 1))
        r.cells[0].width = Inches(2.2)
        r.cells[1].width = Inches(4.8)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Investigation Pipeline Rules Applied
    rules = state.get("applied_rules") or {}
    if rules:
        h1_rules = doc.add_heading("Investigation Pipeline Rules Applied", level=3)
        h1_rules.paragraph_format.space_before = Pt(6)
        h1_rules.paragraph_format.space_after = Pt(4)

        stage_order = [
            "intake_safety", "feature_extraction", "ml_ensemble",
            "explainability", "deep_threat_analysis", "mitre_mapping",
            "analyst_review", "cognitive_triage", "soc_report"
        ]
        active_rules = [(sk.replace("_", " ").title(), str(rules[sk])) for sk in stage_order if sk in rules]

        if active_rules:
            t_rules = doc.add_table(rows=len(active_rules) + 1, cols=2)
            t_rules.style = "Table Grid"
            t_rules.alignment = WD_TABLE_ALIGNMENT.CENTER
            _set_cell_text(t_rules.rows[0].cells[0], "Stage", bold=True, bg_color="0F172A", text_color="FFFFFF")
            _set_cell_text(t_rules.rows[0].cells[1], "Trigger & Operational Rule Evaluation", bold=True, bg_color="0F172A", text_color="FFFFFF")
            t_rules.rows[0].cells[0].width = Inches(2.2)
            t_rules.rows[0].cells[1].width = Inches(4.8)

            for idx, (stitle, rdesc) in enumerate(active_rules):
                r = t_rules.rows[idx + 1]
                _set_cell_text(r.cells[0], stitle, bold=True, bg_color="F8FAFC")
                _set_cell_text(r.cells[1], rdesc)
                r.cells[0].width = Inches(2.2)
                r.cells[1].width = Inches(4.8)

            doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Autonomous Decision Rationale Callout
    deep = state.get("deep_analysis") or {}
    if deep.get("rationale"):
        p_dec = doc.add_paragraph()
        r_dec = p_dec.add_run(f"Autonomous Forensics Decision: {deep['rationale']}")
        r_dec.bold = True
        r_dec.font.size = Pt(9)
        p_dec.paragraph_format.space_after = Pt(6)

    # 2. AI Threat Intelligence Narrative
    llm = state.get("llm_triage") or {}
    raw_narrative = llm.get("executive_threat_narrative") or state.get("llm_narrative", "")
    if raw_narrative:
        h2 = doc.add_heading("2. AI Threat Intelligence & Cognitive Forensics", level=2)
        h2.paragraph_format.space_before = Pt(10)
        h2.paragraph_format.space_after = Pt(4)

        provider_name = llm.get("provider", "ExplainPhish AI")
        model_name = llm.get("model", "Cognitive Reasoner")
        status_badge = llm.get("status_badge", "")
        p_note = doc.add_paragraph()
        if status_badge:
            r_note = p_note.add_run(f"{status_badge} (Engine: {provider_name} | Model: {model_name})")
        else:
            r_note = p_note.add_run(f"Cognitive triage generated by {provider_name} ({model_name}):")
        r_note.italic = True
        r_note.font.size = Pt(8.5)
        p_note.paragraph_format.space_after = Pt(4)

        paragraphs = raw_narrative.split("\n\n")
        for p in paragraphs:
            clean_p = _clean_markdown_text(p)
            if clean_p:
                p_body = doc.add_paragraph()
                p_body.paragraph_format.space_after = Pt(4)
                r_body = p_body.add_run(clean_p)
                r_body.font.size = Pt(9.5)

    # 3. Model Predictions Table
    preds = state.get("model_predictions", [])
    if preds:
        h3 = doc.add_heading("3. Multi-Model ML Ensemble Consensus", level=2)
        h3.paragraph_format.space_before = Pt(10)
        h3.paragraph_format.space_after = Pt(4)

        p_pintro = doc.add_paragraph()
        r_pintro = p_pintro.add_run(f"Consensus voting evaluated across {len(preds)} standardized architectures with 100% feature parity.")
        r_pintro.font.size = Pt(9)
        p_pintro.paragraph_format.space_after = Pt(4)

        t_pred = doc.add_table(rows=len(preds) + 1, cols=4)
        t_pred.style = "Table Grid"
        t_pred.alignment = WD_TABLE_ALIGNMENT.CENTER

        headers = ["Architecture", "Prediction", "Phishing Probability", "Status"]
        col_widths = [Inches(2.5), Inches(2.0), Inches(1.8), Inches(0.7)]
        for j, h in enumerate(headers):
            _set_cell_text(t_pred.rows[0].cells[j], h, bold=True, bg_color="0F172A", text_color="FFFFFF")
            t_pred.rows[0].cells[j].width = col_widths[j]

        for i, p in enumerate(preds):
            row = t_pred.rows[i + 1]
            pl = "MALICIOUS" if p.get("prediction") == 1 else "BENIGN"
            prob_val = p.get("probability", 0.0)
            icon = "⚠️" if p.get("prediction") == 1 else "✅"
            _set_cell_text(row.cells[0], str(p.get("model", "-")), bold=True)
            _set_cell_text(row.cells[1], pl, bold=(pl == "MALICIOUS"))
            _set_cell_text(row.cells[2], f"{prob_val:.1%}")
            _set_cell_text(row.cells[3], icon)
            for j in range(4):
                row.cells[j].width = col_widths[j]

        consensus_text = "Unanimous agreement across all models." if state.get("unanimous") else "Split consensus vote resolved through LangGraph deep analysis."
        p_c = doc.add_paragraph()
        r_c = p_c.add_run(f"Consensus Status: {consensus_text}")
        r_c.font.size = Pt(9)
        r_c.italic = True
        p_c.paragraph_format.space_after = Pt(6)

    # 4. Explainability Decision Drivers
    drivers = state.get("top_risk_drivers", [])
    if drivers:
        h4 = doc.add_heading("4. Explainable AI (XAI) — Top Decision Drivers", level=2)
        h4.paragraph_format.space_before = Pt(10)
        h4.paragraph_format.space_after = Pt(4)

        p_xintro = doc.add_paragraph()
        r_xintro = p_xintro.add_run("Feature impact on model classification derived from normalized feature importance and standardized training deviations (z = (x - μ) / σ):")
        r_xintro.font.size = Pt(9)
        p_xintro.paragraph_format.space_after = Pt(4)

        t_drv = doc.add_table(rows=len(drivers[:6]) + 1, cols=6)
        t_drv.style = "Table Grid"
        t_drv.alignment = WD_TABLE_ALIGNMENT.CENTER

        headers = ["Feature", "Description", "Raw Value", "Z-Score", "Direction", "Risk Impact"]
        col_w = [Inches(1.4), Inches(2.3), Inches(0.8), Inches(0.7), Inches(0.5), Inches(1.3)]
        for j, h in enumerate(headers):
            _set_cell_text(t_drv.rows[0].cells[j], h, bold=True, bg_color="0F172A", text_color="FFFFFF")
            t_drv.rows[0].cells[j].width = col_w[j]

        for i, d in enumerate(drivers[:6]):
            row = t_drv.rows[i + 1]
            is_risk = d.get("direction") == "↑"
            dir_text = "🔺 Increases Risk" if is_risk else "🔻 Reduces Risk"
            _set_cell_text(row.cells[0], str(d.get("feature", "-")), bold=True)
            _set_cell_text(row.cells[1], str(d.get("description", "-")))
            _set_cell_text(row.cells[2], str(d.get("raw_value", "-")))
            _set_cell_text(row.cells[3], f"{d.get('std_value', 0.0):+.2f}")
            _set_cell_text(row.cells[4], str(d.get("direction", "")))
            _set_cell_text(row.cells[5], f"{dir_text} ({d.get('impact', 0.0):+.4f})")
            for j in range(6):
                row.cells[j].width = col_w[j]

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 5. Deep Forensic Static Telemetry & IOCs
    indicators = deep.get("indicators", [])
    iocs = state.get("indicators_of_compromise", [])
    if indicators or iocs:
        h5 = doc.add_heading("5. Deep Forensic Static Telemetry", level=2)
        h5.paragraph_format.space_before = Pt(10)
        h5.paragraph_format.space_after = Pt(4)

        if indicators:
            p_ind_title = doc.add_paragraph()
            r_it = p_ind_title.add_run("Autonomous static inspection findings:")
            r_it.font.size = Pt(9.5)
            r_it.italic = True
            p_ind_title.paragraph_format.space_after = Pt(2)
            for ind in indicators:
                p_ind = doc.add_paragraph(style='List Bullet')
                r_item = p_ind.add_run(ind)
                r_item.font.size = Pt(9)
                p_ind.paragraph_format.space_after = Pt(2)

        if iocs:
            h5_ioc = doc.add_heading("Extracted Indicators of Compromise (IOCs)", level=3)
            h5_ioc.paragraph_format.space_before = Pt(6)
            h5_ioc.paragraph_format.space_after = Pt(4)

            t_ioc = doc.add_table(rows=len(iocs[:10]) + 1, cols=2)
            t_ioc.style = "Table Grid"
            t_ioc.alignment = WD_TABLE_ALIGNMENT.CENTER

            headers = ["Indicator Type", "Extracted Target / IOC Value"]
            for j, h in enumerate(headers):
                _set_cell_text(t_ioc.rows[0].cells[j], h, bold=True, bg_color="0F172A", text_color="FFFFFF")
                t_ioc.rows[0].cells[j].width = Inches(2.5) if j == 0 else Inches(4.5)

            for i, item in enumerate(iocs[:10]):
                row = t_ioc.rows[i + 1]
                _set_cell_text(row.cells[0], str(item.get("type", "-")), bold=True)
                _set_cell_text(row.cells[1], str(item.get("value", "-")))
                row.cells[0].width = Inches(2.5)
                row.cells[1].width = Inches(4.5)

            doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 6. MITRE ATT&CK Mapping
    mitre = state.get("mitre_tactics", [])
    if mitre:
        h6_m = doc.add_heading("6. MITRE ATT&CK Mapping", level=2)
        h6_m.paragraph_format.space_before = Pt(10)
        h6_m.paragraph_format.space_after = Pt(4)

        t_mitre = doc.add_table(rows=len(mitre) + 1, cols=4)
        t_mitre.style = "Table Grid"
        t_mitre.alignment = WD_TABLE_ALIGNMENT.CENTER

        headers = ["Technique ID", "Tactic", "Technique Name", "Operational Description"]
        col_w_m = [Inches(1.2), Inches(1.5), Inches(1.8), Inches(2.5)]
        for j, h in enumerate(headers):
            _set_cell_text(t_mitre.rows[0].cells[j], h, bold=True, bg_color="0F172A", text_color="FFFFFF")
            t_mitre.rows[0].cells[j].width = col_w_m[j]

        for i, m in enumerate(mitre):
            row = t_mitre.rows[i + 1]
            _set_cell_text(row.cells[0], str(m.get("id", "-")), bold=True)
            _set_cell_text(row.cells[1], str(m.get("tactic", "-")))
            _set_cell_text(row.cells[2], str(m.get("technique", "-")), bold=True)
            _set_cell_text(row.cells[3], str(m.get("description", "-")))
            for j in range(4):
                row.cells[j].width = col_w_m[j]

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 7. Prescriptive SOAR Incident Response Playbook
    playbook = state.get("soc_playbook_actions", [])
    if playbook:
        h7_s = doc.add_heading("7. Prescriptive SOAR Incident Response Playbook", level=2)
        h7_s.paragraph_format.space_before = Pt(10)
        h7_s.paragraph_format.space_after = Pt(4)

        t_soar = doc.add_table(rows=len(playbook) + 1, cols=3)
        t_soar.style = "Table Grid"
        t_soar.alignment = WD_TABLE_ALIGNMENT.CENTER

        headers = ["Stage", "Action Tier", "Prescribed Response Action"]
        col_w_s = [Inches(1.6), Inches(1.8), Inches(3.6)]
        for j, h in enumerate(headers):
            _set_cell_text(t_soar.rows[0].cells[j], h, bold=True, bg_color="0F172A", text_color="FFFFFF")
            t_soar.rows[0].cells[j].width = col_w_s[j]

        for i, a in enumerate(playbook):
            row = t_soar.rows[i + 1]
            _set_cell_text(row.cells[0], str(a.get("stage", "-")), bold=True)
            _set_cell_text(row.cells[1], str(a.get("tier", "-")), bold=True)
            _set_cell_text(row.cells[2], str(a.get("action", "-")))
            for j in range(3):
                row.cells[j].width = col_w_s[j]

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 8. Human-in-the-Loop (HITL) Analyst Verification
    analyst_review = state.get("analyst_review") or {}
    if analyst_review:
        h8_h = doc.add_heading("8. Human-in-the-Loop (HITL) Analyst Verification", level=2)
        h8_h.paragraph_format.space_before = Pt(10)
        h8_h.paragraph_format.space_after = Pt(4)

        hitl_rows = [
            ("Review Status", str(analyst_review.get("status", "N/A"))),
            ("Analyst Decision", str(analyst_review.get("decision", "APPROVE"))),
            ("Signed-off By", str(analyst_review.get("analyst_id", "SOC-Analyst"))),
            ("Review Timestamp", str(analyst_review.get("reviewed_at", timestamp))),
            ("Analyst Justification", str(analyst_review.get("notes", "Standard verification sign-off."))),
        ]

        t_hitl = doc.add_table(rows=len(hitl_rows), cols=2)
        t_hitl.style = "Table Grid"
        t_hitl.alignment = WD_TABLE_ALIGNMENT.CENTER

        for idx, (lbl, val) in enumerate(hitl_rows):
            r = t_hitl.rows[idx]
            _set_cell_text(r.cells[0], lbl, bold=True, bg_color="F1F5F9")
            _set_cell_text(r.cells[1], val, bold=(idx == 1))
            r.cells[0].width = Inches(2.2)
            r.cells[1].width = Inches(4.8)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    p_footer = doc.add_paragraph()
    p_footer.paragraph_format.space_before = Pt(14)
    p_footer.paragraph_format.space_after = Pt(4)
    r_foot = p_footer.add_run("Report automatically generated by ExplainPhish Autonomous Security Agent. For SOC escalations, submit hash to SIEM.")
    r_foot.font.size = Pt(8.5)
    r_foot.italic = True
    r_foot.font.color.rgb = RGBColor(100, 116, 139)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _set_cell_text(cell, text: str, bold: bool = False, bg_color: Optional[str] = None, text_color: Optional[str] = None):
    """Utility to style table cells with typography, color, and background shading."""
    cell.text = text
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    for run in p.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8.5)
        run.bold = bold
        if text_color:
            r = int(text_color[0:2], 16)
            g = int(text_color[2:4], 16)
            b = int(text_color[4:6], 16)
            run.font.color.rgb = RGBColor(r, g, b)

    if bg_color:
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading)
